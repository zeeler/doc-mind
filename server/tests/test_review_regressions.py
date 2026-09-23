"""导入、检索、模型协议和认证的回归测试；数据全部位于临时目录。"""

import json
from unittest.mock import MagicMock, patch

import httpx
import pytest
from sqlalchemy import text

from server.config import AppConfig
from server.database import get_engine, get_session_ctx, init_db, reset_engine
from server.models.document import Document, DocumentChunk


@pytest.fixture
def isolated_db(tmp_data_dir, monkeypatch):
    monkeypatch.setattr("server.database.DATA_DIR", tmp_data_dir)
    monkeypatch.setattr("server.services.pipeline.DATA_DIR", tmp_data_dir)
    monkeypatch.setattr("server.routers.documents.DATA_DIR", tmp_data_dir)
    reset_engine()
    AppConfig.invalidate_cache()
    init_db()
    yield tmp_data_dir
    get_engine().dispose()
    reset_engine()
    AppConfig.invalidate_cache()


def add_doc(doc_id="doc", path="/unused.txt"):
    with get_session_ctx() as session:
        session.add(Document(id=doc_id, title="测试文档", file_name=path.split("/")[-1],
                             file_type=path.rsplit(".", 1)[-1], file_path=path))


def fts_rows():
    with get_engine().connect() as conn:
        return conn.execute(text("SELECT chunk_id, content FROM chunks_fts")).all()


def test_pipeline_commits_searchable_fts_and_replaces_old_rows(isolated_db):
    """普通导入/重新索引都必须同步替换片段和 FTS，不能吞掉写锁异常。"""
    from server.services.pipeline import index_document
    from server.services.search import SearchService
    add_doc()
    with patch("server.services.pipeline.VectorStore"):
        index_document("doc", "原来的内容", {})
        index_document("doc", "全文检索哨兵", {})
    with get_session_ctx() as session:
        chunks = session.query(DocumentChunk).all()
        assert len(chunks) == 1
        assert [row[0] for row in fts_rows()] == [chunks[0].id]
    hits = SearchService(isolated_db)._fts_search("全文检索哨兵")
    assert len(hits) == 1
    assert hits[0]["content"] == "全文检索哨兵"


def test_pipeline_fts_failure_rolls_back_chunks(isolated_db):
    from server.services.pipeline import index_document
    add_doc()
    with patch("server.services.pipeline.VectorStore"), patch(
        "server.services.pipeline.fts_insert", side_effect=RuntimeError("FTS write failed")
    ):
        with pytest.raises(RuntimeError, match="FTS write failed"):
            index_document("doc", "不能假装成功", {})
    with get_session_ctx() as session:
        assert session.query(DocumentChunk).count() == 0
        assert session.get(Document, "doc").chunk_count == 0


def test_migration_repairs_missing_fts_for_existing_chunks(isolated_db):
    add_doc()
    with get_session_ctx() as session:
        session.add(DocumentChunk(id="c-old", document_id="doc", chunk_no=1, content="旧索引遗漏"))
    with get_engine().begin() as conn:
        conn.execute(text("PRAGMA user_version = 5"))
    init_db()
    assert [row[0] for row in fts_rows()] == ["c-old"]
    init_db()
    assert len(fts_rows()) == 1


@pytest.mark.parametrize("filename", ["index.md", "INDEX.MD", "normal.md"])
def test_worker_preserves_uploaded_markdown(isolated_db, filename):
    from server.models.job import Job
    from server.services.worker import _execute_job
    path = isolated_db / "files" / "doc" / filename
    path.parent.mkdir()
    original = "# 原始资料\n" + "正文内容。" * 300 + "TAIL_SENTINEL"
    path.write_text(original)
    add_doc(path=str(path))
    with patch("server.services.pipeline.index_document") as index:
        for job_type in ("quick_scan", "full_index"):
            with get_session_ctx() as session:
                job = Job(document_id="doc", job_type=job_type, status="running")
                session.add(job)
                session.commit()
            with patch("server.services.worker.AppConfig.get_all", return_value={"auto_tag_enabled": "false"}):
                _execute_job(job)
            assert path.read_text() == original
        assert index.call_args.args[1] == original


@pytest.mark.parametrize("stream", [False, True])
@pytest.mark.asyncio
async def test_anthropic_system_prompt_is_top_level(stream):
    from server.services.llm import LLMAdapter
    adapter = LLMAdapter({"llm_provider": "claude", "claude_api_key": "dummy"})
    messages = [{"role": "system", "content": "系统规则"},
                {"role": "user", "content": "问题"}]
    bodies = []

    def respond(request):
        body = json.loads(request.content)
        bodies.append(body)
        if stream:
            return httpx.Response(200, text='data: {"type":"content_block_delta","delta":{"type":"text_delta","text":"答案"}}\n\n',
                                  headers={"content-type": "text/event-stream"})
        return httpx.Response(200, json={"content": [{"type": "text", "text": "答案"}]})

    transport = httpx.MockTransport(respond)
    if stream:
        client = httpx.AsyncClient(transport=transport)
        with patch("server.services.llm.httpx.AsyncClient", return_value=client):
            output = [part async for part in adapter.chat_stream(messages)]
            assert output[0]["content"] == "答案"
    else:
        client = httpx.Client(transport=transport)
        with patch("server.services.llm.httpx.Client", return_value=client):
            assert adapter.chat(messages)["content"] == "答案"
    assert bodies[0].get("system") == "系统规则"
    assert bodies[0]["messages"] == [{"role": "user", "content": [{"type": "text", "text": "问题"}]}]


@pytest.mark.parametrize("stream", [False, True])
@pytest.mark.asyncio
async def test_selected_web_search_without_parallel_prefetch(stream):
    from server.services.rag import RAGService
    retriever = MagicMock()
    retriever.retrieve.return_value = []
    config = {"web_search_enabled": "true", "web_search_speculative": "false",
              "tavily_api_key": "dummy", "anysearch_enabled": "false"}
    web_hit = {"chunk_id": "web-1", "document_title": "网络资料", "content": "联网找到的内容",
               "url": "https://example.com", "match_type": "web"}
    with patch("server.services.rag.WebSearchClient.search", return_value=[web_hit]), patch("server.services.rag.LLMAdapter") as llm:
        llm.return_value.chat.return_value = {"content": "答案"}
        async def tokens(**kwargs):
            yield {"type": "token", "content": "答案"}
        llm.return_value.chat_stream = tokens
        rag = RAGService(retriever, config)
        if stream:
            output = [part async for part in rag.ask_stream("问题")]
            citations = next(part["data"] for part in output if part["type"] == "citations")
        else:
            citations = rag.ask_sync("问题")["citations"]
        assert len(citations) == 1
        assert citations[0]["url"] == "https://example.com"


def test_docx_keeps_table_text_in_document_order(tmp_path):
    from docx import Document as WordDocument
    from server.services.parser import parse_file
    word = WordDocument()
    word.add_paragraph("表格前")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "预算"
    table.cell(1, 1).text = "12345"
    nested = table.cell(1, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "嵌套说明"
    word.add_paragraph("表格后")
    path = tmp_path / "table.docx"
    word.save(path)
    parsed = parse_file(path)
    assert "项目" in parsed and "金额" in parsed and "嵌套说明" in parsed
    assert parsed.index("表格前") < parsed.index("12345") < parsed.index("表格后")


@pytest.fixture
def memory_manager(tmp_data_dir):
    from server.services.memory_manager import MemoryManager
    return MemoryManager(config={"memory_export_auto": "false"}, persist_dir=str(tmp_data_dir / "chroma"))


def test_recall_filters_session_before_top_k(memory_manager):
    mgr = memory_manager
    for i in range(20):
        mgr.store.add(f"other-{i}", "临时项目预算", {"scope": "session", "source_conv_id": "A"})
    mgr.store.add("mine", "临时项目预算", {"scope": "session", "source_conv_id": "B"})
    mgr.store.add("global", "长期全局偏好", {"scope": "global"})
    hits = mgr.recall("临时项目预算", conv_id="B", top_k=5)
    assert {hit["id"] for hit in hits} == {"mine", "global"}
    assert {hit["id"] for hit in mgr.recall("预算")} == {"global"}


def test_memory_dedup_preserves_scope_and_conversation(memory_manager):
    mgr = memory_manager
    ids = [mgr.memorize("相同的预算决定", scope=scope, metadata={"source_conv_id": conv})
           for scope, conv in [("global", "A"), ("session", "A"), ("session", "B")]]
    assert len(set(ids)) == 3
    assert mgr.memorize("相同的预算决定", scope="session", metadata={"source_conv_id": "B"}) == ids[2]
    assert mgr.consolidate(dry_run=True)["pairs"] == []


@pytest.fixture
def auth_client(isolated_db):
    from fastapi.testclient import TestClient
    from server.main import app
    AppConfig().set("api_key", "review-test-key")
    return TestClient(app)


def test_browser_login_authenticates_requests_and_preview(auth_client, isolated_db):
    path = isolated_db / "files" / "preview.txt"
    path.write_text("原文预览")
    add_doc(path=str(path))
    assert auth_client.get("/api/v1/documents").status_code == 401
    login = auth_client.post("/api/v1/auth/login", headers={"X-API-Key": "review-test-key"})
    assert login.status_code == 200
    assert "HttpOnly" in login.headers["set-cookie"]
    assert "SameSite=strict" in login.headers["set-cookie"]
    assert auth_client.get("/api/v1/documents/doc/file").text == "原文预览"
    assert auth_client.post("/api/v1/conversations", json={}).status_code == 200
    AppConfig().set("api_key", "rotated-key")
    assert auth_client.get("/api/v1/documents").status_code == 401


def test_browser_auth_rejects_wrong_key_and_cross_origin_cookie(auth_client):
    assert auth_client.post("/api/v1/auth/login", headers={"X-API-Key": "wrong"}).status_code == 401
    assert auth_client.post("/api/v1/auth/login", headers={"X-API-Key": "review-test-key"}).status_code == 200
    response = auth_client.post("/api/v1/conversations", json={}, headers={"Origin": "http://evil.example"})
    assert response.status_code == 403
    assert auth_client.get("/api/v1/conversations", headers={"Origin": "http://evil.example"}).status_code == 403
    assert auth_client.post("/api/v1/conversations", json={}, headers={"Origin": "http://testserver"}).status_code == 200


@pytest.mark.parametrize("cookie", ["9" * 400 + ".abc", "bad-token", "0.invalid"])
def test_malformed_browser_cookie_is_unauthorized(auth_client, cookie):
    auth_client.cookies.set("kb_session", cookie, path="/api/")
    assert auth_client.get("/api/v1/conversations").status_code == 401


def test_browser_session_rejects_expiry_future_and_tampering(auth_client):
    import time
    from server.middleware.auth import SESSION_MAX_AGE, create_session_token
    now = time.time()
    for issued in (now - SESSION_MAX_AGE - 1, now + 60):
        with patch("server.middleware.auth.time.time", return_value=issued):
            token = create_session_token("review-test-key")
        auth_client.cookies.set("kb_session", token, path="/api/")
        assert auth_client.get("/api/v1/conversations").status_code == 401
    token = create_session_token("review-test-key")
    auth_client.cookies.set("kb_session", token + "bad", path="/api/")
    assert auth_client.get("/api/v1/conversations").status_code == 401
    auth_client.cookies.set("kb_session", token, path="/api/")
    assert auth_client.get("/api/v1/conversations").status_code == 200


def test_browser_cookie_authenticates_upload_and_chat_stream(auth_client):
    AppConfig().set("memory_enabled", "false")
    auth_client.post("/api/v1/auth/login", headers={"X-API-Key": "review-test-key"})
    upload = auth_client.post("/api/v1/documents/upload", files={"file": ("sample.txt", b"uploaded content", "text/plain")})
    assert upload.status_code == 200
    doc_id = upload.json()["data"]["id"]
    assert auth_client.get(f"/api/v1/documents/{doc_id}/file").content == b"uploaded content"
    conversation = auth_client.post("/api/v1/conversations", json={}).json()["data"]["id"]

    async def answer(*args, **kwargs):
        yield {"type": "token", "content": "流式认证通过"}
        yield {"type": "citations", "data": []}
    rag = MagicMock()
    rag.ask_stream = answer
    with patch("server.services.registry.ServiceRegistry.get_rag_service", return_value=rag):
        response = auth_client.post("/api/v1/chat/stream", json={"conversation_id": conversation, "question": "测试"})
    assert response.status_code == 200
    assert "流式认证通过" in response.text
    messages = auth_client.get(f"/api/v1/conversations/{conversation}").json()["data"]["messages"]
    assert messages[-1]["content"] == "流式认证通过"


def test_embedding_fallback_keeps_fts_in_sync_after_partial_batch(isolated_db, monkeypatch):
    from server.services.pipeline import index_document
    add_doc()
    monkeypatch.setattr("server.services.pipeline._EMBED_BATCH_SIZE", 1)
    with patch("server.services.pipeline.VectorStore"), patch("server.services.pipeline.Embedder") as embedder, patch(
        "server.services.pipeline.chunk_text", return_value=["第一个片段", "第二个片段"]
    ):
        embedder.return_value.embed.side_effect = [[[0.1, 0.2]], [[0.1, 0.2]], RuntimeError("embedding interrupted")]
        index_document("doc", "任意输入", {"embedding_enabled": "true"})
    with get_session_ctx() as session:
        chunks = session.query(DocumentChunk).all()
        assert len(chunks) == 2
        assert session.get(Document, "doc").chunk_count == 2
        assert sorted(row[0] for row in fts_rows()) == sorted(chunk.id for chunk in chunks)


def test_api_header_and_query_key_still_work(auth_client):
    assert auth_client.get("/api/v1/conversations", headers={"X-API-Key": "review-test-key"}).status_code == 200
    assert auth_client.get("/api/v1/conversations", params={"api_key": "review-test-key"}).status_code == 200
    assert auth_client.get("/api/v1/conversations", headers={"X-API-Key": "wrong"}).status_code == 401
